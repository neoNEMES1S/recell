"""Synthetic format fixtures test inspection, not battery research results."""

import json
from datetime import datetime
from pathlib import Path

import h5py
import numpy as np
import pytest
import yaml
from openpyxl import Workbook
from scipy.io import savemat

from second_life_battery.data.inventory import apply_review, build_manifest, main
from second_life_battery.utils.config import load_config


def test_empty_and_missing_dataset(tmp_path: Path) -> None:
    assert build_manifest(tmp_path)["status"] == "empty_dataset"
    assert build_manifest(tmp_path / "missing")["status"] == "missing_directory"


def test_workbook_preamble_multiple_sheets_and_reviewed_header(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "measurements"
    sheet.append(["instrument export"])
    sheet.append([])
    sheet.append(["when", "unusual_capacity_label", "temperature"])
    sheet.append([datetime(2024, 1, 1), 4.1, None])
    sheet.append([datetime(2024, 1, 2), 4.0, 25])
    book.create_sheet("empty")
    book.save(tmp_path / "arbitrary.xlsx")
    unreviewed = build_manifest(tmp_path)
    record = unreviewed["files"][0]
    assert record["cell_id"] is None
    assert record["tables"][0]["number_of_observations"] is None
    assert record["tables"][0]["candidate_headers"][1]["row"] == 3
    reviewed = build_manifest(tmp_path, header_rows={"arbitrary.xlsx": {"measurements": 3}})
    table = reviewed["files"][0]["tables"][0]
    assert table["number_of_observations"] == 2
    assert table["detected_columns"][1] == "unusual_capacity_label"
    assert table["missing_values_by_column_position"]["3"] == 1
    assert table["native_datetime_ranges_by_column_position"]["1"] == [
        "2024-01-01T00:00:00",
        "2024-01-02T00:00:00",
    ]
    assert reviewed["targets_verified"] is False


def test_csv_preserves_duplicate_headers_and_missing_values(tmp_path: Path) -> None:
    (tmp_path / "data.csv").write_text("A,A,B\n1,,3\n2,4,\n", encoding="utf-8")
    table = build_manifest(tmp_path, header_rows={"data.csv": {"table": 1}})["files"][0]["tables"][
        0
    ]
    assert table["detected_columns"] == ["A", "A", "B"]
    assert table["missing_values_by_column_position"] == {"1": 0, "2": 1, "3": 1}


def test_mat_formats_preserve_shapes_without_guessing_observations(tmp_path: Path) -> None:
    savemat(tmp_path / "legacy.mat", {"trace": np.zeros((7, 3)), "meta": {"name": "A"}})
    with h5py.File(tmp_path / "new.mat", "w") as handle:
        handle.create_dataset("cell/trace", data=np.ones((3, 7)))
    files = build_manifest(tmp_path)["files"]
    assert files[0]["variables"][0]["shape"] == [7, 3]
    assert files[1]["variables"][0]["name"] == "cell/trace"
    assert all(file["number_of_observations"] is None for file in files)


def test_bad_files_are_recorded_and_other_files_survive(tmp_path: Path) -> None:
    (tmp_path / "bad.xlsx").write_bytes(b"not a zip file")
    (tmp_path / "bad.mat").write_bytes(b"bad mat file")
    (tmp_path / "readme.txt").write_text("metadata")
    (tmp_path / "valid.tsv").write_text("x\ty\n1\t2\n")
    result = build_manifest(tmp_path)
    assert result["file_count"] == 4
    assert result["status"] == "inspection_errors"
    assert result["status_counts"] == {"error": 2, "unsupported": 1, "inspected": 1}
    assert all(len(file["sha256"]) == 64 for file in result["files"])


def test_bad_header_mapping_is_not_silently_ignored(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="absent files"):
        build_manifest(tmp_path, header_rows={"typo.xlsx": {"sheet": 1}})
    (tmp_path / "x.csv").write_text("a,b\n1,2\n")
    result = build_manifest(tmp_path, header_rows={"x.csv": {"table": 20}})
    assert result["status"] == "inspection_errors"


def test_cli_writes_empty_manifest_and_resolved_config(tmp_path: Path) -> None:
    configs = tmp_path / "configs"
    configs.mkdir()
    config = configs / "base.yaml"
    config.write_text(
        yaml.safe_dump(
            {
                "dataset": {"raw_dir": "data/raw"},
                "inspection": {"output_dir": "outputs/inventory"},
            }
        )
    )
    assert main(["--config", str(config)]) == 2
    output = tmp_path / "outputs/inventory"
    manifest = json.loads((output / "dataset_manifest.json").read_text())
    assert manifest["status"] == "missing_directory"
    assert manifest["files"] == []
    assert (output / "files.csv").read_text().startswith("source_file,")
    assert json.loads((output / "inspection.jsonl").read_text())["level"] == "INFO"
    snapshot = yaml.safe_load((output / "config.snapshot.yaml").read_text())
    assert snapshot["dataset"]["raw_dir"] == str(tmp_path / "data/raw")
    assert (output / "environment.json").exists()


def test_cli_prevents_raw_output_contamination(tmp_path: Path) -> None:
    with pytest.raises(SystemExit) as error:
        main(["--raw-dir", str(tmp_path), "--output-dir", str(tmp_path / "out")])
    assert error.value.code == 2


@pytest.mark.parametrize("rows", [0, -1, 101, True, "20"])
def test_invalid_preview_config(tmp_path: Path, rows: object) -> None:
    path = tmp_path / "config.yaml"
    path.write_text(
        yaml.safe_dump(
            {
                "dataset": {"raw_dir": "raw"},
                "inspection": {"output_dir": "out", "preview_rows": rows},
            }
        )
    )
    with pytest.raises(ValueError, match="preview_rows"):
        load_config(path)


def test_sampled_counts_are_never_reported_as_complete(tmp_path: Path) -> None:
    (tmp_path / "data.csv").write_text("a,b\n1,2\n3,4\n5,6\n")
    result = build_manifest(
        tmp_path, preview_rows=2, scan_limit=2, header_rows={"data.csv": {"table": 1}}
    )
    table = result["files"][0]["tables"][0]
    assert table["scan_complete"] is False
    assert table["number_of_observations"] is None
    assert table["sample_observations"] == 1


def test_review_preserves_conflicting_identity_labels() -> None:
    manifest = {
        "files": [
            {"source_file": "cycling_tests/G1/G1.xlsx", "extension": ".xlsx", "cell_id": None},
            {
                "source_file": "diagnostic_tests/W8/scan_w9.xlsx",
                "extension": ".xlsx",
                "cell_id": None,
            },
        ]
    }
    apply_review(
        manifest,
        {
            "source": "review.md",
            "cells": {
                "G1": {"duty_cycle_type": "commercial"},
                "W8": {"duty_cycle_type": "commercial"},
                "W9": {"duty_cycle_type": "residential"},
            },
            "test_type_rules": [{"pattern": "cycling_tests/*", "test_type": "cycling"}],
        },
    )
    assert manifest["battery_ids"] == ["G1"]
    assert manifest["files"][0]["duty_cycle_type"] == "commercial"
    assert manifest["files"][1]["cell_id"] is None
    assert manifest["files"][1]["identity_status"] == "ambiguous"


def test_parallel_inspection_retains_all_files_and_order(tmp_path: Path) -> None:
    (tmp_path / "b.csv").write_text("x\n1\n")
    (tmp_path / "a.csv").write_text("x\n2\n")
    serial = build_manifest(tmp_path)
    parallel = build_manifest(tmp_path, workers=2)
    assert serial["files"] == parallel["files"]


def test_reviewed_headers_file_loads_into_snapshot_config(tmp_path: Path) -> None:
    headers = {"x.xlsx": {"Sheet1": 1}}
    (tmp_path / "headers.json").write_text(json.dumps(headers))
    path = tmp_path / "config.yaml"
    path.write_text(
        yaml.safe_dump(
            {
                "dataset": {"raw_dir": "raw"},
                "inspection": {
                    "output_dir": "out",
                    "header_rows_file": "headers.json",
                },
            }
        )
    )
    config = load_config(path)
    assert config["inspection"]["header_rows"] == headers
    assert "header_rows_file" not in config["inspection"]


def test_review_assigns_only_complete_native_timestamp_coverage(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Channel_1"
    sheet.append(["Date_Time", "Capacity"])
    sheet.append([datetime(2024, 1, 1), 4.2])
    sheet.append([datetime(2024, 1, 2), 4.1])
    book.save(tmp_path / "G1.xlsx")
    review = {
        "source": "review.md",
        "cells": {"G1": {"duty_cycle_type": "commercial"}},
        "measurement_sheet_prefix": "Channel_",
        "timestamp_column": "Date_Time",
    }
    for limit, expected_count in [(None, 2), (2, None)]:
        manifest = build_manifest(
            tmp_path, preview_rows=2, scan_limit=limit, header_rows={"G1.xlsx": {"Channel_1": 1}}
        )
        apply_review(manifest, review)
        record = manifest["files"][0]
        assert record["number_of_observations"] == expected_count
        assert (record["date_time_coverage"] is None) == (limit is not None)
