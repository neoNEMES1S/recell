"""Inspect raw file structure without inferring battery identities or scientific units."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import platform
import re
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, timezone
from fnmatch import fnmatchcase
from importlib.metadata import version
from pathlib import Path
from typing import Any, Iterable
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import h5py
import openpyxl
import yaml
from scipy.io import loadmat, whosmat
from scipy.io.matlab import MatlabOpaque, MatReadError

from second_life_battery.utils.config import load_config
from second_life_battery.utils.logging import configure_logging

LOGGER = logging.getLogger("second_life_battery")
UNVERIFIED = ["cell_id", "test_type", "duty_cycle_type", "timestamp", "capacity", "units"]


def sha256_file(path: Path) -> str:
    """Hash a source in bounded memory for reproducible provenance."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (value != value or abs(value) == float("inf")):
        return str(value)
    return value


def inspect_rows(
    rows: Iterable[tuple[Any, ...]],
    name: str,
    preview_rows: int,
    header_row: int | None,
    scan_limit: int | None = None,
) -> dict[str, Any]:
    """Stream rows; preserve preambles and distinguish candidate from reviewed headers."""
    preview = []
    candidates = []
    columns = None
    nonempty_rows = 0
    observations = 0
    missing: Counter[int] = Counter()
    present: Counter[int] = Counter()
    date_ranges: dict[str, list[str]] = {}
    max_columns = 0
    scan_complete = True
    for row_number, row in enumerate(rows, start=1):
        if scan_limit is not None and row_number > scan_limit:
            scan_complete = False
            break
        populated = [value for value in row if value is not None and value != ""]
        if not populated:
            continue
        nonempty_rows += 1
        width = max(i + 1 for i, value in enumerate(row) if value is not None and value != "")
        max_columns = max(max_columns, width)
        values = [_json_value(value) for value in row[:width]]
        if len(preview) < preview_rows:
            preview.append({"row": row_number, "values": values})
            if all(isinstance(value, str) for value in populated):
                candidates.append({"row": row_number, "columns": values})
        if row_number == header_row:
            columns = values
        if header_row is not None and row_number > header_row:
            observations += 1
            for index, value in enumerate(row):
                if value is not None and value != "":
                    present[index] += 1
        for index, value in enumerate(row):
            if isinstance(value, (datetime, date)):
                stamp = value.isoformat()
                key = str(index + 1)
                lo, hi = date_ranges.get(key, [stamp, stamp])
                date_ranges[key] = [min(lo, stamp), max(hi, stamp)]
    if header_row is not None:
        if columns is None:
            raise ValueError(f"Header row {header_row} is empty or absent in {name}")
        missing.update({index: observations - present[index] for index in range(max_columns)})
    return {
        "name": name,
        "scan_complete": scan_complete,
        "nonempty_rows_including_preamble": nonempty_rows,
        "max_populated_columns": max_columns,
        "preview": preview,
        "candidate_headers": candidates,
        "header_row": header_row,
        "detected_columns": columns,
        "number_of_observations": observations
        if header_row is not None and scan_complete
        else None,
        "sample_observations": observations if header_row is not None else None,
        "observation_count_definition": "nonempty rows after reviewed header; may include footers",
        "missing_values_by_column_position": {str(i + 1): n for i, n in missing.items()},
        "native_datetime_ranges_by_column_position": date_ranges,
        "date_time_coverage": None,
        "warnings": [
            "Counts, missingness and datetime ranges cover scanned rows only; check scan_complete.",
            "Header candidates are unverified; numeric and text timestamps are not interpreted.",
            "Native datetime ranges describe cell values, not verified experiment timestamps.",
        ],
    }


def inspect_mat(path: Path) -> dict[str, Any]:
    """Read MATLAB variable metadata without loading large measurement arrays."""
    variables: list[dict[str, Any]] = []
    if h5py.is_hdf5(path):
        with h5py.File(path, "r") as handle:

            def visit(name: str, obj: h5py.Group | h5py.Dataset) -> None:
                if isinstance(obj, h5py.Dataset):
                    variables.append(
                        {
                            "name": name,
                            "shape": list(obj.shape),
                            "dtype": str(obj.dtype),
                        }
                    )

            handle.visititems(visit)
        file_format = "hdf5_mat_candidate"
    else:
        try:
            variables = [
                {"name": name, "shape": list(shape), "matlab_class": kind}
                for name, shape, kind in whosmat(str(path))
            ]
        except TypeError:
            # scipy whosmat cannot report the opaque MATLAB tables in this release.
            if path.stat().st_size > 16 * 1024 * 1024:
                raise
            contents = loadmat(path)
            if not any(isinstance(value, MatlabOpaque) for value in contents.values()):
                raise
            variables = [
                {
                    "name": name,
                    "shape": list(value.shape),
                    "matlab_class": type(value).__name__,
                    "dtype": str(value.dtype),
                }
                for name, value in contents.items()
                if not name.startswith("__")
            ]
        file_format = "mat_v4_to_v7_2"
    return {
        "format": file_format,
        "variables": variables,
        "warnings": [
            "Variable shapes are not observation counts. MATLAB structs and HDF5 references "
            "require a semantic parser after review; metadata inspection does not decode them."
        ],
    }


def inspect_file(
    path: Path,
    root: Path,
    preview_rows: int,
    headers: dict[str, int],
    scan_limit: int | None = None,
) -> dict[str, Any]:
    """Inventory one file, recording expected read/format failures without dropping the file."""
    relative = path.relative_to(root).as_posix()
    record: dict[str, Any] = {
        "source_file": relative,
        "extension": path.suffix.lower(),
        "status": "pending",
        "cell_id": None,
        "test_type": None,
        "duty_cycle_type": None,
        "date_time_coverage": None,
        "number_of_observations": None,
        "detected_columns": None,
        "missing_fields": None,
        "unverified_fields": UNVERIFIED.copy(),
        "tables": [],
    }
    try:
        if path.is_symlink():
            record.update(status="skipped_symlink", reason="Point raw_dir at the real directory")
            return record
        record["size_bytes"] = path.stat().st_size
        record["sha256"] = sha256_file(path)
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                unknown = set(headers) - set(workbook.sheetnames)
                if unknown:
                    raise ValueError(f"Configured sheets absent from workbook: {sorted(unknown)}")
                for sheet in workbook:
                    declared_rows, declared_columns = sheet.max_row, sheet.max_column
                    # Scan real cells even when an exporter writes incorrect sheet dimensions.
                    sheet.reset_dimensions()
                    record["tables"].append(
                        inspect_rows(
                            sheet.iter_rows(values_only=True),
                            sheet.title,
                            preview_rows,
                            headers.get(sheet.title),
                            scan_limit,
                        )
                    )
                    record["tables"][-1].update(
                        declared_max_row=declared_rows,
                        declared_max_column=declared_columns,
                    )
            finally:
                workbook.close()
        elif suffix in {".csv", ".tsv"}:
            if set(headers) - {"table"}:
                raise ValueError("CSV/TSV header mapping must use table name 'table'")
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ",", strict=True)
                record["tables"].append(
                    inspect_rows(
                        (tuple(row) for row in rows),
                        "table",
                        preview_rows,
                        headers.get("table"),
                        scan_limit,
                    )
                )
        elif suffix == ".mat":
            record.update(inspect_mat(path))
        else:
            record.update(status="unsupported", reason="File listed and hashed; content not parsed")
            return record
        record["status"] = "inspected"
    except (
        OSError,
        ValueError,
        TypeError,
        BadZipFile,
        ParseError,
        csv.Error,
        NotImplementedError,
        EOFError,
        KeyError,
        MatReadError,
    ) as error:
        record.update(status="error", error_type=type(error).__name__, error=str(error))
        LOGGER.error("file_inspection_failed: %s", error, extra={"source_file": relative})
    return record


def build_manifest(
    raw_dir: Path,
    preview_rows: int = 20,
    header_rows: dict[str, dict[str, int]] | None = None,
    scan_limit: int | None = None,
    workers: int = 1,
) -> dict[str, Any]:
    """Inspect files recursively; missing data produces an explicit incomplete manifest."""
    if not 1 <= preview_rows <= 100:
        raise ValueError("preview_rows must be between 1 and 100")
    if scan_limit is not None and (type(scan_limit) is not int or scan_limit < preview_rows):
        raise ValueError("scan_limit must be null or an integer at least preview_rows")
    root = raw_dir.resolve()
    if root.exists() and not root.is_dir():
        raise ValueError(f"Raw path is not a directory: {root}")
    paths = sorted(path for path in root.rglob("*") if path.is_file() and path.name != ".gitkeep")
    mappings = header_rows or {}
    unused = set(mappings) - {p.relative_to(root).as_posix() for p in paths}
    if unused:
        raise ValueError(f"Header mappings reference absent files: {sorted(unused)}")
    arguments = [
        (path, root, preview_rows, mappings.get(path.relative_to(root).as_posix(), {}), scan_limit)
        for path in paths
    ]
    if workers == 1:
        records = [inspect_file(*args) for args in arguments]
    else:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            records = list(pool.map(_inspect_file_args, arguments))
    for record in records:
        LOGGER.log(
            logging.ERROR if record["status"] == "error" else logging.INFO,
            "file_inspected: %s",
            record["status"],
            extra={"source_file": record["source_file"]},
        )
    counts = Counter(record["status"] for record in records)
    if not root.exists():
        status = "missing_directory"
    elif not records:
        status = "empty_dataset"
    elif counts["error"]:
        status = "inspection_errors"
    elif not counts["inspected"]:
        status = "no_supported_data"
    else:
        status = "inspected_pending_semantic_review"
    return {
        "schema_version": 1,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "raw_dir": str(root),
        "status": status,
        "battery_ids": [],
        "battery_ids_status": "unverified",
        "targets_verified": False,
        "file_count": len(records),
        "status_counts": dict(counts),
        "extension_counts": dict(Counter(r["extension"] for r in records)),
        "files": records,
        "limitations": [
            "No battery IDs, units, duty cycles, capacity targets or test types are inferred.",
            "missing_fields=null means unknown, not that all required fields are present.",
            "Sheet counts must not be summed as independent battery observations.",
            "Archives and unsupported formats require separate review or extraction.",
        ],
    }


def _inspect_file_args(arguments: tuple) -> dict[str, Any]:
    """Multiprocessing entry point; files are independent and results retain sorted order."""
    return inspect_file(*arguments)


def apply_review(manifest: dict[str, Any], review: dict[str, Any]) -> None:
    """Attach configured, evidence-backed path identities; retain conflicts for review."""
    cells = review.get("cells", {})
    if not cells:
        return
    pattern = re.compile(
        r"(?<![a-z0-9])(" + "|".join(map(re.escape, cells)) + r")(?![a-z0-9])", re.IGNORECASE
    )
    discovered = set()
    for record in manifest["files"]:
        if record["extension"] not in {".xlsx", ".mat"}:
            continue
        candidates = sorted({m.group().upper() for m in pattern.finditer(record["source_file"])})
        record["cell_id_candidates"] = candidates
        record["quality_flags"] = []
        record["identity_status"] = "ambiguous" if len(candidates) > 1 else "unassigned"
        if len(candidates) > 1:
            record["quality_flags"].append("conflicting_cell_labels")
        if record.get("tables") and not any(
            table["nonempty_rows_including_preamble"] for table in record["tables"]
        ):
            record["quality_flags"].append("empty_workbook")
        if len(candidates) == 1:
            cell = candidates[0]
            record.update(
                cell_id=cell,
                duty_cycle_type=cells[cell]["duty_cycle_type"],
                identity_status="path_label_requires_measurement_validation",
                identity_evidence=review["source"],
            )
            discovered.add(cell)
        for rule in review.get("test_type_rules", []):
            if fnmatchcase(record["source_file"], rule["pattern"]):
                record["test_type"] = rule["test_type"]
                break
        prefix = review.get("measurement_sheet_prefix")
        tables = [
            table
            for table in record.get("tables", [])
            if prefix and table["name"].startswith(prefix)
        ]
        if len(tables) == 1:
            table = tables[0]
            record["observation_table"] = table["name"]
            record["number_of_observations"] = table["number_of_observations"]
            record["detected_columns"] = table["detected_columns"]
            columns = table["detected_columns"] or []
            time_column = review.get("timestamp_column")
            if time_column in columns and table["scan_complete"]:
                position = str(columns.index(time_column) + 1)
                record["date_time_coverage"] = table[
                    "native_datetime_ranges_by_column_position"
                ].get(position)
                record["date_time_basis"] = "native Excel datetime; timezone unverified"
    manifest["battery_ids"] = sorted(discovered)
    manifest["battery_ids_status"] = "observed_path_labels_matched_to_reviewed_registry"
    manifest["review_source"] = review["source"]


def write_outputs(manifest: dict[str, Any], output_dir: Path, config: dict[str, Any]) -> None:
    """Write JSON provenance, file-level CSV and a human-readable inspection report."""
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "dataset_manifest.json").write_text(
        json.dumps(manifest, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    (output_dir / "config.snapshot.yaml").write_text(yaml.safe_dump(config, sort_keys=False))
    environment = {
        "python": platform.python_version(),
        "platform": platform.platform(),
        "packages": {
            name: version(name)
            for name in ["numpy", "pandas", "scipy", "openpyxl", "h5py", "PyYAML"]
        },
    }
    (output_dir / "environment.json").write_text(json.dumps(environment, indent=2) + "\n")
    fields = [
        "source_file",
        "extension",
        "size_bytes",
        "sha256",
        "status",
        "cell_id",
        "test_type",
        "duty_cycle_type",
        "number_of_observations",
        "error",
    ]
    with (output_dir / "files.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(manifest["files"])
    availability = []
    for record in manifest["files"]:
        for table in record["tables"]:
            columns = table["detected_columns"] or []
            for index, column in enumerate(columns, start=1):
                availability.append(
                    {
                        "source_file": record["source_file"],
                        "table": table["name"],
                        "column_position": index,
                        "column": column,
                        "scan_complete": table["scan_complete"],
                        "scanned_observations": table["sample_observations"],
                        "missing_values": table["missing_values_by_column_position"].get(
                            str(index)
                        ),
                    }
                )
    with (output_dir / "column_availability.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=[
                "source_file",
                "table",
                "column_position",
                "column",
                "scan_complete",
                "scanned_observations",
                "missing_values",
            ],
        )
        writer.writeheader()
        writer.writerows(availability)
    complete_tables = sum(
        table["scan_complete"] for record in manifest["files"] for table in record["tables"]
    )
    report = [
        "# Dataset inspection",
        "",
        f"Status: **{manifest['status']}**",
        "",
        f"Raw directory: `{manifest['raw_dir']}`",
        f"Files found: {manifest['file_count']}",
        f"Observed battery labels: {', '.join(manifest['battery_ids']) or 'unverified'}",
        f"Fully scanned tables: {complete_tables}",
        "",
        "Target verification: **pending**. Training is not implemented.",
        "",
        "## File summary",
        "",
        "| Extension | Files |",
        "| --- | ---: |",
        *[
            f"| {ext or '(none)'} | {count} |"
            for ext, count in manifest["extension_counts"].items()
        ],
        "",
        "## Limitations",
        "",
        *[f"- {item}" for item in manifest["limitations"]],
        "",
        "## Next review",
        "",
        "Inspect sheet previews / MATLAB metadata in dataset_manifest.json. Confirm cell IDs,",
        "header rows, time origins, units, RPT discharge segments and duplicated raw/processed",
        "sources before building a parser. Populate reviewed header_rows in the config and rerun.",
    ]
    (output_dir / "inspection_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    """Run inspection; exit 2 for missing/empty data and 1 for file read failures."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=Path("configs/base.yaml"))
    parser.add_argument(
        "--raw-dir", type=Path, help="Override raw path, relative to current directory"
    )
    parser.add_argument("--output-dir", type=Path, help="Override output path")
    parser.add_argument("--full-scan", action="store_true", help="Scan every workbook/CSV row")
    args = parser.parse_args(argv)
    config = load_config(args.config)
    project_root = args.config.resolve().parent.parent
    raw = (args.raw_dir or project_root / config["dataset"]["raw_dir"]).resolve()
    output = (args.output_dir or project_root / config["inspection"]["output_dir"]).resolve()
    if output == raw or output.is_relative_to(raw):
        parser.error("Output directory must be outside the raw directory")
    output.mkdir(parents=True, exist_ok=True)
    logger = configure_logging(output / "inspection.jsonl")
    config["dataset"]["raw_dir"] = str(raw)
    config["inspection"]["output_dir"] = str(output)
    if args.full_scan:
        config["inspection"]["scan_limit"] = None
    manifest = build_manifest(
        raw,
        config["inspection"].get("preview_rows", 20),
        config["inspection"].get("header_rows", {}),
        config["inspection"].get("scan_limit"),
        config["inspection"].get("workers", 1),
    )
    apply_review(manifest, config.get("review", {}))
    write_outputs(manifest, output, config)
    logger.info(
        "inspection_complete: status=%s files=%s", manifest["status"], manifest["file_count"]
    )
    if manifest["status"] == "inspection_errors":
        return 1
    return 0 if manifest["status"] == "inspected_pending_semantic_review" else 2


if __name__ == "__main__":
    raise SystemExit(main())
